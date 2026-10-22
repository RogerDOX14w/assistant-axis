"""SOC 2018 codes, the BLS OEWS national table and the broad groups (coding plan 16, § 5b and § 7d).

**Codes.**  An O*NET-SOC code (``29-1212.00``) folds onto its 2018 SOC detailed code (``29-1212``), and that onto its
broad group (``29-1210``): the code with its last digit set to 0, **except** where a broad group has more than nine
detailed occupations and its codes run on into the next ten: SOC 2018's Physicians, ``29-1210``, whose detailed
codes run from 29-1211 to 29-1229 (there is no broad group 29-1220).  :data:`BROAD_SPILLOVER` holds that case; with
it the 867 detailed codes of O*NET 31.0 fold onto the 459 broad groups of SOC 2018 (460 without it).

**OEWS.**  :func:`read_oews_national` reads the May 2025 national estimates (``oesm25nat.zip``: an xlsx inside a
zip, or the xlsx itself) with ``openpyxl``, imported only there (it is not a dependency yet: Roger approved adding
it with the downloads, Decisions for Roger 6); :func:`read_oews_table` reads the same columns from a CSV or JSON
file (tests, or a hand-made table).  ``OCC_CODE -> {"title", "o_group", "tot_emp"}``; ``tot_emp`` is None where BLS
suppresses it (``**``, ``*``, ``#``).  Column names are matched without regard to case.

**Broad groups** (:func:`broad_groups`): one per broad code of the O*NET occupations, with its O*NET occupations
(every O*NET-SOC row under it, ``.00`` and the ``.01`` ... specialties alike, in code order), its title (OEWS's
broad row, else the SOC structure's, else the single detailed occupation's own title, else the O*NET titles joined,
``title_source`` saying which) and its employment (OEWS's broad row, else the sum of its detailed rows,
``employment_source`` saying which; None with no OEWS).
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable, Mapping, Optional

from .onet import Occupation

#: Detailed-code prefixes (``XX-XXX``) whose broad group is not the code with a final 0 (module docstring).
BROAD_SPILLOVER: dict[str, str] = {"29-122": "29-1210"}
#: SOC 2018 minor groups whose code is not ``XX-X000`` (the rule's exceptions as far as known; used only by
#: :func:`soc_tiers` when no OEWS table says otherwise).
MINOR_EXCEPTIONS: dict[str, str] = {"15-12": "15-1200", "31-11": "31-1100", "51-51": "51-5100"}
_CODE_RE = re.compile(r"^(\d{2})-(\d{4})(?:\.(\d{2}))?$")
SUPPRESSED = ("**", "*", "#", "~", "")
OEWS_COLUMNS = ("occ_code", "occ_title", "o_group", "tot_emp")


class MissingDependency(ImportError):
    """``openpyxl`` is needed to read a BLS workbook."""


def _check(code: str) -> re.Match:
    m = _CODE_RE.match(str(code).strip())
    if not m:
        raise ValueError(f"not a SOC or O*NET-SOC code: {code!r}")
    return m


def detailed_code(code: str) -> str:
    """``29-1212.00`` -> ``29-1212`` (a SOC code is returned as it is)."""
    m = _check(code)
    return f"{m.group(1)}-{m.group(2)}"


def broad_code(onet_or_soc_code: str) -> str:
    """``29-1212.00`` -> ``29-1210``; ``29-1221.00`` -> ``29-1210`` (:data:`BROAD_SPILLOVER`); ``29-1210`` -> itself."""
    d = detailed_code(onet_or_soc_code)
    if d[:6] in BROAD_SPILLOVER:
        return BROAD_SPILLOVER[d[:6]]
    return d[:6] + "0"


def soc_tiers(code: str, *, oews: Optional[Mapping[str, Mapping]] = None) -> dict:
    """``{"major", "minor", "broad", "detailed"}`` of a code (``detailed`` None for a broad code).  The minor group
    is the OEWS table's minor row when one matches, else :data:`MINOR_EXCEPTIONS`, else ``XX-X000``."""
    d = detailed_code(code)
    major = d[:2] + "-0000"
    minor = None
    if oews:
        for cand in (d[:5] + "00", d[:4] + "000"):
            r = oews.get(cand)
            if r and str(r.get("o_group") or "").lower() == "minor":
                minor = cand
                break
    if minor is None:
        minor = MINOR_EXCEPTIONS.get(d[:5], d[:4] + "000")
    broad = broad_code(d)
    return {"major": major, "minor": minor, "broad": broad, "detailed": None if d == broad else d}


# --------------------------------------------------------------------------- OEWS

def _emp(value) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip().replace(",", "")
    if s in SUPPRESSED:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def _rows_to_table(header: list, rows: Iterable) -> dict[str, dict]:
    low = [str(h or "").strip().lower() for h in header]
    idx = {c: low.index(c) for c in OEWS_COLUMNS if c in low}
    missing = [c for c in OEWS_COLUMNS if c not in idx]
    if missing:
        raise ValueError(f"OEWS table without the columns {missing} (found {header})")
    out: dict[str, dict] = {}
    for r in rows:
        r = list(r)
        if len(r) <= max(idx.values()):
            continue
        code = str(r[idx["occ_code"]] or "").strip()
        if not code:
            continue
        out.setdefault(code, {"title": str(r[idx["occ_title"]] or "").strip(),
                              "o_group": str(r[idx["o_group"]] or "").strip().lower(),
                              "tot_emp": _emp(r[idx["tot_emp"]])})
    return out


def _openpyxl():
    try:
        import openpyxl  # noqa: F401 - imported here only (Decisions for Roger 6)
    except ImportError as exc:
        raise MissingDependency(
            "reading a BLS workbook needs openpyxl, which is not installed: run `uv add openpyxl` (approved by Roger "
            "with the downloads, coding_plan_16 Decisions for Roger 6), or give a CSV or JSON table with the columns "
            "OCC_CODE, OCC_TITLE, O_GROUP, TOT_EMP (read_oews_table)") from exc
    return openpyxl


def _xlsx_rows(data: bytes) -> tuple[list, list]:
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header: list = []
        for row in it:                      # the header is the first row naming OCC_CODE
            if row and any(str(c or "").strip().lower() == "occ_code" for c in row):
                header = list(row)
                break
        rows = [list(r) for r in it]
    finally:
        wb.close()
    return header, rows


def read_oews_national(zip_path: Path) -> dict[str, dict]:
    """``OCC_CODE -> {"title", "o_group", "tot_emp"}`` from the OEWS national file: a zip holding the national
    xlsx (the first ``.xlsx`` member whose name has ``nat``, else the first ``.xlsx``), or the xlsx itself.  Needs
    ``openpyxl`` (:class:`MissingDependency` names ``uv add openpyxl`` when it is missing)."""
    p = Path(zip_path)
    if p.suffix.lower() == ".zip":
        with zipfile.ZipFile(p) as zf:
            names = [n for n in zf.namelist() if n.lower().endswith(".xlsx") and not n.split("/")[-1].startswith("~")]
            if not names:
                raise ValueError(f"{p}: no .xlsx inside")
            pick = next((n for n in names if "nat" in n.split("/")[-1].lower()), names[0])
            data = zf.read(pick)
    else:
        data = p.read_bytes()
    header, rows = _xlsx_rows(data)
    if not header:
        raise ValueError(f"{p}: no header row naming OCC_CODE")
    return _rows_to_table(header, rows)


def read_oews_table(path: Path) -> dict[str, dict]:
    """The same table from a CSV (header row) or JSON file (a list of row objects, or ``{code: {title, o_group,
    tot_emp}}``)."""
    p = Path(path)
    if p.suffix.lower() == ".json":
        d = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(d, Mapping):
            return {str(k): {"title": v.get("title"), "o_group": str(v.get("o_group") or "").lower(),
                             "tot_emp": _emp(v.get("tot_emp"))} for k, v in d.items()}
        header = list(d[0].keys()) if d else list(OEWS_COLUMNS)
        return _rows_to_table(header, [[r.get(h) for h in header] for r in d])
    with open(p, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return {}
    return _rows_to_table(rows[0], rows[1:])


def read_oews(path: Path) -> dict[str, dict]:
    """:func:`read_oews_national` for a ``.zip`` or ``.xlsx``, :func:`read_oews_table` for ``.csv`` / ``.json``."""
    suf = Path(path).suffix.lower()
    if suf in (".zip", ".xlsx"):
        return read_oews_national(path)
    if suf in (".csv", ".json"):
        return read_oews_table(path)
    raise ValueError(f"{path}: an OEWS table is .zip, .xlsx, .csv or .json")


def read_soc_structure(path: Path) -> dict[str, str]:
    """``code -> title`` from the BLS SOC 2018 structure file (optional; xlsx through ``openpyxl``, or CSV): every
    row whose first code-shaped cell is a SOC code, its title the next non-empty text cell.  The xlsx layout was
    not seen when this was written; the reader looks for the codes rather than for named columns."""
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        openpyxl = _openpyxl()
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        with open(p, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
    out: dict[str, str] = {}
    for r in rows:
        cells = [str(c).strip() if c is not None else "" for c in r]
        for i, c in enumerate(cells):
            if _CODE_RE.match(c) and "." not in c:
                title = next((x for x in cells[i + 1:] if x and not _CODE_RE.match(x)), None)
                if title:
                    out.setdefault(c, title)
                break
    return out


# --------------------------------------------------------------------------- broad groups

@dataclass
class BroadGroup:
    code: str
    title: str
    detailed: list = field(default_factory=list)          # list[Occupation]
    employment: Optional[int] = None
    title_source: str = "oews"
    employment_source: Optional[str] = "oews_broad"

    def as_table_row(self) -> dict:
        """``table.jsonl``'s broad row (coding plan 16, § 6.6)."""
        return {"kind": "broad", "code": self.code, "title": self.title, "employment": self.employment,
                "title_source": self.title_source, "employment_source": self.employment_source,
                "detailed": [asdict(o) for o in self.detailed]}


def broad_groups(onet_rows: Iterable[Occupation], oews: Optional[Mapping[str, Mapping]] = None, *,
                 soc_titles: Optional[Mapping[str, str]] = None) -> list[BroadGroup]:
    """One :class:`BroadGroup` per broad code of ``onet_rows``, in code order (module docstring)."""
    by: dict[str, list[Occupation]] = {}
    for o in onet_rows:
        by.setdefault(broad_code(o.code), []).append(o)
    oews = oews or {}
    out = []
    for code in sorted(by):
        occ = sorted(by[code], key=lambda o: o.code)
        row = oews.get(code)
        if row and row.get("title"):
            title, tsrc = str(row["title"]), "oews"
        elif soc_titles and soc_titles.get(code):
            title, tsrc = soc_titles[code], "soc_structure"
        else:
            dets = sorted({detailed_code(o.code) for o in occ})
            main = [o for o in occ if o.code.endswith(".00")]
            if len(dets) == 1 and main:
                title, tsrc = main[0].title, "onet_detailed"
            else:
                title, tsrc = "; ".join(o.title for o in (main or occ)[:3]) + ("; ..." if len(main or occ) > 3 else ""), \
                    "onet_titles"
        emp, esrc = None, None
        if row and row.get("tot_emp") is not None:
            emp, esrc = int(row["tot_emp"]), "oews_broad"
        elif oews:
            vals = [oews[d]["tot_emp"] for d in sorted({detailed_code(o.code) for o in occ})
                    if d in oews and oews[d].get("tot_emp") is not None]
            if vals:
                emp, esrc = int(sum(vals)), "oews_detailed_sum"
        out.append(BroadGroup(code=code, title=title, detailed=occ, employment=emp, title_source=tsrc,
                              employment_source=esrc))
    return out


def detailed_employment(onet_code: str, oews: Optional[Mapping[str, Mapping]], groups: Mapping[str, BroadGroup]
                        ) -> tuple[Optional[int], Optional[str]]:
    """The employment a job title under ``onet_code`` is ranked by: its detailed SOC code's OEWS row, else its broad
    group's (plan § 7d: a job title takes its occupation's detailed employment)."""
    d = detailed_code(onet_code)
    if oews and d in oews and oews[d].get("tot_emp") is not None:
        return int(oews[d]["tot_emp"]), "oews_detailed"
    g = groups.get(broad_code(onet_code))
    if g is not None and g.employment is not None:
        return g.employment, "broad_group"
    return None, None


__all__ = ["BROAD_SPILLOVER", "MINOR_EXCEPTIONS", "MissingDependency", "detailed_code", "broad_code", "soc_tiers",
           "read_oews_national", "read_oews_table", "read_oews", "read_soc_structure", "BroadGroup", "broad_groups",
           "detailed_employment"]
